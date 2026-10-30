#!/usr/bin/env python3
import sys
import os
import json
import time

# T35 (t35-design D2, D4, D5, D6): OCR for a page whose text layer is EMPTY — the pure part (the trigger, the thread
# pin, the reading order, the page-entry fields) lives beside this script in ocr_text.py; the engine (RapidOCR on the
# CPU, the image's bundled models, onnxruntime) is constructed lazily, ONCE per bridge process, only when a page needs
# it — a document with text on every page never loads it.
# No bytecode cache for the imported module: a `__pycache__/ocr_text.cpython-*.pyc` written under src/ at the first
# ingest moved the pod's own tree-hash computation off the build's (I-HEALTH on the after chain — §11a A14; F-153).
sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ocr_text

# T18 (design D10): the fallback branch (a PDF pypdfium2 could not open) renders its FIRST N pages for the captioner —
# one number, never a list of page numbers from one document.
FALLBACK_RENDER_PAGES = int(os.environ.get("NEXCODE_FALLBACK_RENDER_PAGES", "15"))

_OCR_ENGINE = {"engine": None, "error": None, "label": None, "threads": None, "tried": False}


def _ocr_engine():
    """RapidOCR with the intra-op threads pinned to the container's CPU quota (F-148: the engine's default is every
    core the node shows — 30–45× slower on a 2-CPU limit, t35-design §0a), the bundled PP-OCRv6 small models, CPU.
    A construction failure (no onnxruntime — F-127's shape) is recorded once; every OCR page is then counted as
    failed and named, and the file keeps its captions (a failed pass never fails the file)."""
    if _OCR_ENGINE["tried"]:
        return _OCR_ENGINE
    _OCR_ENGINE["tried"] = True
    try:
        import hashlib
        import importlib.metadata
        import rapidocr as rapidocr_pkg
        from rapidocr import RapidOCR
        threads = ocr_text.cpu_quota_threads()
        engine = RapidOCR(params=ocr_text.engine_params(threads))
        models = {}
        mdir = os.path.join(os.path.dirname(rapidocr_pkg.__file__), "models")
        if os.path.isdir(mdir):
            for name in sorted(os.listdir(mdir)):
                if name.endswith(".onnx"):
                    with open(os.path.join(mdir, name), "rb") as f:
                        models[name] = hashlib.md5(f.read()).hexdigest()
        try:
            import onnxruntime
            ort_version = getattr(onnxruntime, "__version__", "?")
        except Exception:
            ort_version = "?"
        try:
            rapid_version = importlib.metadata.version("rapidocr")
        except Exception:
            rapid_version = getattr(rapidocr_pkg, "__version__", "?")
        _OCR_ENGINE.update({"engine": engine, "threads": threads, "label": ocr_text.engine_label(rapid_version, ort_version, models, threads)})
    except Exception as err:
        _OCR_ENGINE["error"] = f"{type(err).__name__}: {str(err)[:200]}"
    return _OCR_ENGINE


def new_ocr_stats(mode):
    """What the bridge reports about the OCR pass (D10): the parser copies it onto the trace, the route onto the ledger."""
    return {"mode": mode, "noText": 0, "pages": 0, "empty": 0, "failed": 0, "chars": 0, "ms": 0, "noTextPages": [], "pagesList": [], "emptyPages": [], "failedPages": [], "engine": None, "threads": None, "scale": ocr_text.OCR_SCALE, "error": None}


def ocr_page(page, page_no, stats):
    """OCR one pypdfium2 page (D5): rendered at OCR_SCALE in memory (no file, no crop-folder write — every caption
    memo key stands), read by the engine (rapidocr converts the PIL render to BGR itself, as it does a file), the boxes
    ordered into lines. Returns (text or None, the page entry's fields). A page OCR finds nothing on stays a page
    without text (ocrEmpty, named); a failure never fails the file (ocrFailed, named)."""
    eng = _ocr_engine()
    if eng["engine"] is None:
        stats["failed"] += 1
        stats["failedPages"].append(page_no)
        if stats["error"] is None:
            stats["error"] = eng["error"]
        return None, {"ocr": False, "ocrError": eng["error"]}
    t0 = time.time()
    try:
        img = page.render(scale=ocr_text.OCR_SCALE).to_pil()
        out = eng["engine"](img)
        boxes = [] if out is None or out.boxes is None else [[[float(x), float(y)] for x, y in b] for b in out.boxes]
        txts = [] if out is None or out.txts is None else list(out.txts)
        scores = [] if out is None or out.scores is None else [float(v) for v in out.scores]
        lines = ocr_text.reading_order(boxes, txts, scores)
        ms = (time.time() - t0) * 1000
        stats["ms"] += int(round(ms))
        if not lines:
            stats["empty"] += 1
            stats["emptyPages"].append(page_no)
            return None, {"ocr": False, "ocrEmpty": True, "ocrMs": int(round(ms))}
        fields = ocr_text.page_ocr_entry(lines, len(boxes), ms, min(scores) if scores else None)
        stats["pages"] += 1
        stats["pagesList"].append(page_no)
        stats["chars"] += fields["ocrChars"]
        return ocr_text.lines_text(lines), fields
    except Exception as err:
        stats["failed"] += 1
        stats["failedPages"].append(page_no)
        msg = f"{type(err).__name__}: {str(err)[:200]}"
        if stats["error"] is None:
            stats["error"] = msg
        return None, {"ocr": False, "ocrError": msg}


def finish_ocr_stats(stats):
    eng = _OCR_ENGINE
    if eng["tried"]:
        stats["engine"] = eng["label"]
        stats["threads"] = eng["threads"]
        if stats["error"] is None and eng["error"] is not None:
            stats["error"] = eng["error"]
    return stats


def page_count(doc_path):
    """`--page-count`: the page count of a PDF, read in milliseconds (the parser sizes the bridge's timeout by it — T35 §2.2)."""
    if not os.path.exists(doc_path):
        return json.dumps({"error": f"File not found: {doc_path}"})
    if not doc_path.lower().endswith(".pdf"):
        return json.dumps({"status": "success", "pages": None})
    try:
        import pypdfium2 as pdfium
        return json.dumps({"status": "success", "pages": len(pdfium.PdfDocument(doc_path))})
    except Exception as err:
        return json.dumps({"error": f"page count failed: {err}"})


def pptx_slide_meta(doc_path):
    """T41 (design D6, D7): per slide (1-based, the slide's position in the file — the page number docling gives it,
    hidden slides included): whether it is hidden (`show="0"`), and the lines of its FOOTER / SLIDE_NUMBER / DATE
    placeholders (by placeholder TYPE — never by text). A deck python-pptx cannot read raises; the caller says so."""
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    chrome_types = (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE)
    out = {}
    prs = Presentation(doc_path)
    for idx, slide in enumerate(prs.slides, 1):
        chrome = set()
        for shape in slide.placeholders:
            try:
                ptype = shape.placeholder_format.type
            except Exception:
                ptype = None
            if ptype in chrome_types and getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for line in shape.text_frame.text.split("\n"):
                    if line.strip():
                        chrome.add(line.strip())
        out[idx] = {"hidden": slide.element.get("show") == "0", "chrome": chrome}
    return out


def convert_document_to_markdown_pages(doc_path, output_dir=None, no_images=False, ocr_mode="auto"):
    if not os.path.exists(doc_path):
        return json.dumps({"error": f"File not found: {doc_path}"})
    # T35 (design D3): the OCR mode the parser passed (`--ocr auto | off`); every return path reports the pass.
    ocr_mode = ocr_text.normalise_mode(ocr_mode)
    ocr_stats = new_ocr_stats(ocr_mode)

    # T24 (design D6, D7): `--no-images` = text only. No page render, no
    # figure export, no crop folder created, `pictures` empty on every path.
    # Without it every path behaves byte-for-byte as before.
    if no_images:
        output_dir = None
    
    # For PDF documents, use pypdfium2 for ultra-fast, zero-OOM memory-safe parsing & figure rendering
    if doc_path.lower().endswith(".pdf"):
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(doc_path)
            pages_list = []
            pictures_list = []
            
            if output_dir:
                os.makedirs(output_dir, exist_ok=True)

            for page_idx, page in enumerate(pdf):
                page_no = page_idx + 1
                text = page.get_textpage().get_text_range()
                page_entry = {
                    "page": page_no,
                    "markdown": text
                }
                # T35 (design D2, D5, D6): a page whose text layer is EMPTY is read by OCR under mode auto — the OCR
                # lines become its markdown (the parser chunks them as any page's text, marked textSource ocr), and
                # the page entry says so. Under off such a page is counted and named, and stays without text.
                if not ocr_text.has_text_layer(text):
                    ocr_stats["noText"] += 1
                    ocr_stats["noTextPages"].append(page_no)
                    if ocr_text.page_needs_ocr(text, ocr_mode):
                        ocr_txt, fields = ocr_page(page, page_no, ocr_stats)
                        if ocr_txt is not None:
                            text = ocr_txt
                            page_entry["markdown"] = text
                        page_entry.update(fields)
                pages_list.append(page_entry)
                
                # Adaptive Scale Rendering: Detect Site Plans, Maps, or Detailed Blueprints for Scale=3.5 High-Res Render
                lower_text = text.lower()
                is_dense_map = any(kw in lower_text for kw in ["site plan", "master plan", "phase", "plot", "legend", "layout", "blueprint"])
                render_scale = 3.5 if is_dense_map else 2.0

                # Render every document page image for 100% complete VLM visual coverage
                if output_dir:
                    img = page.render(scale=render_scale).to_pil()
                    filename = f"fig_{page_no}_0.png"
                    abs_path = os.path.join(output_dir, filename)
                    img.save(abs_path, format="PNG")
                    
                    rel_parts = abs_path.split("public")
                    web_rel_path = rel_parts[-1] if len(rel_parts) > 1 else abs_path

                    pictures_list.append({
                        "id": f"fig_{page_no}_0",
                        "page": page_no,
                        "abs_path": abs_path,
                        "web_path": web_rel_path.replace("\\", "/"),
                        "width": img.width,
                        "height": img.height
                    })

            return json.dumps({
                "status": "success",
                "pages": pages_list,
                "pictures": pictures_list,
                "engine": "PyPDFium2 (Fast Multimodal Engine)",
                # True per-page structure: each entry is one physical PDF page.
                "paginated": True,
                # T35 (design D10): the OCR pass — the mode, the pages without a text layer, the pages read / read
                # empty / failed on (named), the characters, the wall time, the engine and its models.
                "ocr": finish_ocr_stats(ocr_stats)
            })
        except Exception as pdf_err:
            sys.stderr.write(f"[PyPDFium2 Error]: {pdf_err}. Falling back to Docling...\n")

    try:
        from docling.document_converter import DocumentConverter, PdfFormatOption
        from docling.datamodel.pipeline_options import PdfPipelineOptions

        pipeline_options = PdfPipelineOptions()
        pipeline_options.do_ocr = False  # Fast native PDF parsing
        pipeline_options.do_table_structure = True  # AI Table grid extraction
        pipeline_options.generate_picture_images = not no_images  # Extract figure crops (T24: off with --no-images)
        pipeline_options.images_scale = 1.0  # Standard scale to maintain light memory footprint

        converter = DocumentConverter(
            format_options={
                'pdf': PdfFormatOption(pipeline_options=pipeline_options)
            }
        )
        
        result = converter.convert(doc_path)
        pages_list = []
        pictures_list = []

        if hasattr(result.document, "pages") and result.document.pages:
            paginated = True
            for page in result.document.pages.values():
                try:
                    # T41 (design D5, F-146): the markdown as the document holds it — no HTML entities (`&amp;`), no
                    # `<!-- image -->` placeholder lines; both were docling export options the bridge never set.
                    page_md = result.document.export_to_markdown(page_no=page.page_no, escape_html=False, image_placeholder="")
                except Exception:
                    page_md = ""
                pages_list.append({
                    "page": page.page_no,
                    "markdown": page_md
                })
        else:
            # No page structure from docling (e.g. .docx has no fixed
            # pagination): the whole document is emitted as a single entry
            # stamped page 1. paginated=False records that this attribution is
            # synthetic, not real (F-11) — the consumer must not present
            # "page 1" as a located citation for this file.
            paginated = False
            # T41 (design D5, F-146): the same two options on the whole-document export.
            full_md = result.document.export_to_markdown(escape_html=False, image_placeholder="")
            pages_list.append({
                "page": 1,
                "markdown": full_md
            })
        
        # T41 (design D6, D7; F-145, F-165): a deck's structure, read with python-pptx (docling's own pptx dependency):
        # a HIDDEN slide (`show="0"`) is not shown to anyone — its page entry carries `hidden` and no markdown, and its
        # pictures are not exported (no chunk, no caption); a slide's FOOTER / SLIDE_NUMBER / DATE placeholder text is
        # not slide content — those lines are removed from that slide's markdown (docling emits them first, which is how
        # the parser took the footer for the slide's heading). Keyed on the placeholder TYPE, never on the text.
        pptx_hidden = set()
        pptx_chrome_removed = 0
        if doc_path.lower().endswith(".pptx"):
            try:
                pptx_meta = pptx_slide_meta(doc_path)
                for entry in pages_list:
                    meta = pptx_meta.get(entry["page"])
                    if not meta:
                        continue
                    if meta["hidden"]:
                        pptx_hidden.add(entry["page"])
                        entry["hidden"] = True
                        entry["markdown"] = ""
                        continue
                    if meta["chrome"]:
                        kept = []
                        for line in entry["markdown"].split("\n"):
                            if line.strip() in meta["chrome"]:
                                pptx_chrome_removed += 1
                                continue
                            kept.append(line)
                        entry["markdown"] = "\n".join(kept)
            except Exception as pptx_err:
                sys.stderr.write(f"[pptx structure]: {pptx_err}; the deck is indexed as docling reads it\n")

        # Save extracted pictures: the largest figure of every page. T31 (F-123):
        # the bridge no longer caps the count — the parser's captionPageBudget
        # is the only bound, and the pages beyond it are counted and named.
        if output_dir and hasattr(result.document, "pictures") and result.document.pictures:
            os.makedirs(output_dir, exist_ok=True)
            page_pics = {}
            
            for idx, pic in enumerate(result.document.pictures):
                try:
                    img = pic.get_image(result.document)
                    if not img:
                        continue
                    # Skip decorative icons / tiny logos (< 180x180 px)
                    if img.width < 180 or img.height < 180:
                        continue
                    
                    page_no = pic.prov[0].page_no if (hasattr(pic, "prov") and pic.prov) else 1
                    # T41 (design D7): a hidden slide's pictures are not exported.
                    if page_no in pptx_hidden:
                        continue
                    fig_id = f"fig_{page_no}_{idx}"
                    
                    if page_no not in page_pics:
                        page_pics[page_no] = []
                    page_pics[page_no].append({
                        "id": fig_id,
                        "page": page_no,
                        "img": img,
                        "area": img.width * img.height
                    })
                except Exception:
                    pass

            # Select the largest primary figure of each page — every page (T31)
            selected_items = []
            for p_no in sorted(page_pics.keys()):
                p_items = page_pics[p_no]
                p_items.sort(key=lambda item: item["area"], reverse=True)
                if p_items:
                    selected_items.append(p_items[0])

            for item in selected_items:
                try:
                    filename = f"{item['id']}.png"
                    abs_path = os.path.join(output_dir, filename)
                    item["img"].save(abs_path, format="PNG")
                    
                    rel_parts = abs_path.split("public")
                    web_rel_path = rel_parts[-1] if len(rel_parts) > 1 else abs_path

                    pictures_list.append({
                        "id": item["id"],
                        "page": item["page"],
                        "abs_path": abs_path,
                        "web_path": web_rel_path.replace("\\", "/"),
                        "width": item["img"].width,
                        "height": item["img"].height
                    })
                except Exception as pic_err:
                    sys.stderr.write(f"[Picture Export Error]: {pic_err}\n")

            import gc
            gc.collect()

        return json.dumps({
            "status": "success",
            "pages": pages_list,
            "page_count": len(pages_list),
            "pictures": pictures_list,
            # This branch previously reported no engine at all, which is how
            # parser.js came to hardcode a label instead (E03). Every return
            # path now names the engine that actually ran.
            "engine": "Docling (Layout AI & Table Transformer)",
            "paginated": paginated,
            # T41 (design D6, D7): what the pptx pass did — the hidden slides (named) and the chrome lines removed.
            "pptx": {"hidden": sorted(pptx_hidden), "chromeLinesRemoved": pptx_chrome_removed} if doc_path.lower().endswith(".pptx") else None,
            # T35: OCR runs only in the pypdfium2 branch (design D5); the pass is reported as not taken here.
            "ocr": finish_ocr_stats(ocr_stats)
        })
    except Exception as err:
        sys.stderr.write(f"[Docling Error]: {err}. Attempting pypdfium2 lightweight fallback...\n")
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(doc_path)
            pages_list = []
            pictures_list = []
            
            if output_dir:
                os.makedirs(output_dir, exist_ok=True)

            for page_idx, page in enumerate(pdf):
                page_no = page_idx + 1
                text = page.get_textpage().get_text_range()
                pages_list.append({
                    "page": page_no,
                    "markdown": text
                })
                
                # Render the first FALLBACK_RENDER_PAGES pages for VLM captioning (T18 D10)
                if output_dir and page_idx < FALLBACK_RENDER_PAGES:
                    img = page.render(scale=1.5).to_pil()
                    filename = f"fig_{page_no}_0.png"
                    abs_path = os.path.join(output_dir, filename)
                    img.save(abs_path, format="PNG")
                    
                    rel_parts = abs_path.split("public")
                    web_rel_path = rel_parts[-1] if len(rel_parts) > 1 else abs_path

                    pictures_list.append({
                        "id": f"fig_{page_no}_0",
                        "page": page_no,
                        "abs_path": abs_path,
                        "web_path": web_rel_path.replace("\\", "/"),
                        "width": img.width,
                        "height": img.height
                    })

            return json.dumps({
                "status": "success",
                "pages": pages_list,
                "pictures": pictures_list,
                "engine": "pypdfium2 (Lightweight Fallback)",
                "paginated": True,
                "ocr": finish_ocr_stats(ocr_stats)
            })
        except Exception as fb_err:
            return json.dumps({"error": f"Docling and PyPDFium2 both failed: {fb_err}"})

def text_word_boxes(page):
    """T45 (design §2.1): every whitespace-delimited word of the page's text layer with its box in points from the
    page's TOP-left — the union of pypdfium2's character boxes over the same text page T18's running-line pass reads
    (`get_text_range`). [] for a page without a text layer. Deterministic for the PDF's bytes and pypdfium2's version
    (5.12.1 in the image); the Site Plan's page 1 gives 567 words. The parser's tile-edge rule (detail_pass.js
    `tileVotes`) reads these against the tiles' boxes: a tile votes on a label only when it holds a whole cluster of
    the label's occurrences. No render, no subprocess, no file; a failure here never fails the page (the caller
    merges without geometry, as before T45)."""
    tp = page.get_textpage()
    n = tp.count_chars()
    if n <= 0:
        return []
    text = tp.get_text_range(0, n)
    # A code point outside the BMP is two UTF-16 units to pdfium and one to Python: read char by char then.
    chars = list(text) if len(text) == n else [tp.get_text_range(k, 1) for k in range(n)]
    _, h = page.get_size()
    words = []
    i = 0
    while i < n:
        if chars[i].isspace() or chars[i] == "":
            i += 1
            continue
        j = i
        while j < n and not chars[j].isspace() and chars[j] != "":
            j += 1
        boxes = [tp.get_charbox(k) for k in range(i, j)]
        l = min(b[0] for b in boxes); r = max(b[2] for b in boxes); bt = min(b[1] for b in boxes); tt = max(b[3] for b in boxes)
        words.append({"text": "".join(chars[i:j]), "x0": round(l, 3), "y0": round(h - tt, 3), "x1": round(r, 3), "y1": round(h - bt, 3)})
        i = j
    return words


def tile_boxes(width, height, rows, cols, overlap):
    """T33 (design D2): the tile geometry — a rows x cols grid over a page of width x height (points for a PDF
    page, pixels for a picture), each tile widened by `overlap` x the page's size on every inner edge. Each box
    is (x0, y0, x1, y1) from the page's TOP-left; `crop` is the amount pypdfium2's render(crop=...) cuts off
    (left, bottom, right, top), PDF y running upward. The same arithmetic as detail_pass.js's tilePlan (pinned
    by U-DETAIL against the design probe's recorded crops); the parser records what this reports."""
    ovx, ovy = width * overlap, height * overlap
    out = []
    for r in range(rows):
        for c in range(cols):
            x0 = max(0.0, c * width / cols - ovx); x1 = min(width, (c + 1) * width / cols + ovx)
            y0 = max(0.0, r * height / rows - ovy); y1 = min(height, (r + 1) * height / rows + ovy)
            out.append({"tile": f"r{r}c{c}", "row": r, "col": c, "box": [x0, y0, x1, y1], "crop": (x0, height - y1, width - x1, y0)})
    return out


def render_detail_tiles(doc_path, output_dir, page_no, rows, cols, overlap, scale, image_path=None):
    """T33 (design D2, D9): the tiles of ONE page for the detail passes. A PDF page is rendered from the PDF —
    the whole page at `scale` and each tile by pypdfium2's crop, so the tile's resolution is independent of the
    page render's keyword rule (F-125) and the existing fig_<page>_0.png is never rewritten (every caption memo
    key stands). A picture from the docling path (a docx / pptx page has no PDF page: `image_path`) is copied
    whole and cut with PIL at its own resolution. Files: detail_<page>_whole.png and detail_<page>_r<r>c<c>.png
    beside the page render — new names, nothing overwritten. Never touches any other path of this script.
    T45 (design §2.1): the reply also carries `textBoxes` — the page's text layer word by word with its box in points
    (text_word_boxes) on the PDF path; None on the picture path (no PDF page); [] for a page without a text layer.
    The tiles' bytes are untouched (every memo key stands)."""
    if not os.path.exists(doc_path):
        return json.dumps({"error": f"File not found: {doc_path}"})
    os.makedirs(output_dir, exist_ok=True)

    def web(abs_path):
        rel_parts = abs_path.split("public")
        return (rel_parts[-1] if len(rel_parts) > 1 else abs_path).replace("\\", "/")

    def entry(fid, abs_path, w, h, extra=None):
        e = {"id": fid, "page": page_no, "abs_path": abs_path, "web_path": web(abs_path), "width": w, "height": h}
        if extra:
            e.update(extra)
        return e

    try:
        whole_path = os.path.join(output_dir, f"detail_{page_no}_whole.png")
        tiles = []
        if image_path is None and doc_path.lower().endswith(".pdf"):
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(doc_path)
            if page_no < 1 or page_no > len(pdf):
                return json.dumps({"error": f"page {page_no} is not in the document ({len(pdf)} page(s))"})
            page = pdf[page_no - 1]
            w, h = page.get_size()
            try:
                text_boxes = text_word_boxes(page)
            except Exception as tb_err:
                sys.stderr.write(f"[tiles] text boxes not read for page {page_no}: {tb_err}; the page merges without geometry\n")
                text_boxes = []
            whole = page.render(scale=scale).to_pil()
            whole.save(whole_path, format="PNG")
            for t in tile_boxes(w, h, rows, cols, overlap):
                img = page.render(scale=scale, crop=t["crop"]).to_pil()
                fn = os.path.join(output_dir, f"detail_{page_no}_{t['tile']}.png")
                img.save(fn, format="PNG")
                tiles.append(entry(f"detail_{page_no}_{t['tile']}", fn, img.width, img.height, {"tile": t["tile"], "row": t["row"], "col": t["col"], "cropPt": [round(v, 3) for v in t["crop"]], "box": [round(v, 3) for v in t["box"]]}))
            return json.dumps({"status": "success", "mode": "tiles", "source": "pdf", "page": page_no, "pagePt": [w, h], "px": [whole.width, whole.height], "scale": scale, "grid": [rows, cols], "overlap": overlap, "whole": entry(f"detail_{page_no}_whole", whole_path, whole.width, whole.height), "tiles": tiles, "textBoxes": text_boxes})
        # A picture (the docling path), or a non-PDF: the image at its own resolution.
        from PIL import Image
        import shutil
        src = image_path if image_path else doc_path
        if not os.path.exists(src):
            return json.dumps({"error": f"tile source not found: {src}"})
        shutil.copyfile(src, whole_path)
        img = Image.open(whole_path)
        w, h = img.width, img.height
        for t in tile_boxes(w, h, rows, cols, overlap):
            x0, y0, x1, y1 = t["box"]
            crop = img.crop((int(x0), int(y0), int(round(x1)), int(round(y1))))
            fn = os.path.join(output_dir, f"detail_{page_no}_{t['tile']}.png")
            crop.save(fn, format="PNG")
            tiles.append(entry(f"detail_{page_no}_{t['tile']}", fn, crop.width, crop.height, {"tile": t["tile"], "row": t["row"], "col": t["col"], "cropPt": None, "box": [int(x0), int(y0), int(round(x1)), int(round(y1))]}))
        return json.dumps({"status": "success", "mode": "tiles", "source": "image", "page": page_no, "pagePt": None, "px": [w, h], "scale": None, "grid": [rows, cols], "overlap": overlap, "whole": entry(f"detail_{page_no}_whole", whole_path, w, h), "tiles": tiles, "textBoxes": None})
    except Exception as err:
        return json.dumps({"error": f"tile rendering failed: {err}"})


def parse_geometry_spec(spec):
    """T46 (D15): `<page>[:<scale>:<ink>,<ink>…]` → (page, scale, inks); the parser passes the pinned setting
    (dimension_chain.js CHAIN_GEOMETRY: 4.0, 200 and 128). Raises on a malformed value."""
    parts = spec.split(":")
    page = int(parts[0])
    scale = float(parts[1]) if len(parts) > 1 and parts[1] else 4.0
    inks = tuple(int(v) for v in parts[2].split(",")) if len(parts) > 2 and parts[2] else (200, 128)
    if page < 1 or scale <= 0 or not inks or any(i < 1 or i > 255 for i in inks):
        raise ValueError(f"geometry spec out of range: {spec}")
    return page, scale, inks


def read_geometry(doc_path, page_no, scale, inks):
    """T46 (design D14 a, D15; dimension_geometry.py): the pixel geometry of ONE PDF page — the page rendered at `scale`
    in memory, its drawing regions' crosshairs and dimension claims read at each ink threshold. Nothing is written
    (no render file, no crop, no memo); a non-PDF has no page to render (an error the parser records — the chain stays
    as read). The module is imported here only, so no other mode loads OpenCV."""
    if not os.path.exists(doc_path):
        return json.dumps({"error": f"File not found: {doc_path}"})
    if not doc_path.lower().endswith(".pdf"):
        return json.dumps({"error": "the geometry reads a PDF page only"})
    try:
        import dimension_geometry
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(doc_path)
        pages = len(pdf)
        pdf.close()
        if page_no < 1 or page_no > pages:
            return json.dumps({"error": f"page {page_no} is not in the document ({pages} page(s))"})
        res = dimension_geometry.read_page(doc_path, page_no, scale, inks)
        return json.dumps({"status": "success", "mode": "geometry", "page": page_no, **res})
    except Exception as err:
        return json.dumps({"error": f"geometry failed: {type(err).__name__}: {str(err)[:300]}"})


def read_colours(doc_path, page_no):
    """T57 (design §2.1, D3, D5; colour_layer.py): the COLOUR LAYER of ONE PDF page for the detail pass's colour check —
    the words (text_word_boxes, the tiles mode's own `textBoxes`) with their font sizes, every filled path with its exact
    fill and outline, and each word's pixel counts per path fill at the tiles' scale. Nothing is written (no render file,
    no crop, no memo); a non-PDF has no page to read (an error the parser records — the page merges as before T57). The
    module is imported here only."""
    if not os.path.exists(doc_path):
        return json.dumps({"error": f"File not found: {doc_path}"})
    if not doc_path.lower().endswith(".pdf"):
        return json.dumps({"error": "the colour layer reads a PDF page only"})
    try:
        import colour_layer
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(doc_path)
        if page_no < 1 or page_no > len(pdf):
            return json.dumps({"error": f"page {page_no} is not in the document ({len(pdf)} page(s))"})
        page = pdf[page_no - 1]
        w, h = page.get_size()
        layer = colour_layer.read_layer(page, text_word_boxes)
        return json.dumps({"status": "success", "mode": "colours", "page": page_no, "pagePt": [w, h], **layer})
    except Exception as err:
        return json.dumps({"error": f"colour layer failed: {type(err).__name__}: {str(err)[:300]}"})


def parse_tiles_spec(spec):
    """`<page>:<rows>x<cols>:<overlap>:<scale>` → (page, rows, cols, overlap, scale); raises on a malformed value."""
    page_s, grid_s, overlap_s, scale_s = spec.split(":")
    rows_s, cols_s = grid_s.lower().split("x")
    page, rows, cols, overlap, scale = int(page_s), int(rows_s), int(cols_s), float(overlap_s), float(scale_s)
    if page < 1 or rows < 1 or cols < 1 or overlap < 0 or overlap >= 0.5 or scale <= 0:
        raise ValueError(f"tiles spec out of range: {spec}")
    return page, rows, cols, overlap, scale


if __name__ == "__main__":
    argv = sys.argv[1:]
    # T33: the tiles mode — its own branch, so every other path stays byte-identical.
    if "--tiles" in argv:
        try:
            spec = argv[argv.index("--tiles") + 1]
            image_path = argv[argv.index("--tile-image") + 1] if "--tile-image" in argv else None
            positional = [a for i, a in enumerate(argv) if not a.startswith("--") and not (i > 0 and argv[i - 1] in ("--tiles", "--tile-image"))]
            if len(positional) < 2:
                raise ValueError("Usage: docling_bridge.py <doc_path> <output_dir> --tiles <page>:<rows>x<cols>:<overlap>:<scale> [--tile-image <picture_abs_path>]")
            page, rows, cols, overlap, scale = parse_tiles_spec(spec)
        except Exception as e:
            print(json.dumps({"error": f"tiles mode: {e}"}))
            sys.exit(1)
        print(render_detail_tiles(positional[0], positional[1], page, rows, cols, overlap, scale, image_path))
        sys.exit(0)

    # T46 (design D15): the geometry mode — its own branch, so every other path stays byte-identical.
    if "--geometry" in argv:
        try:
            spec = argv[argv.index("--geometry") + 1]
            positional = [a for i, a in enumerate(argv) if not a.startswith("--") and not (i > 0 and argv[i - 1] == "--geometry")]
            if len(positional) < 1:
                raise ValueError("Usage: docling_bridge.py <doc_path> --geometry <page>[:<scale>:<ink>,<ink>]")
            page, scale, inks = parse_geometry_spec(spec)
        except Exception as e:
            print(json.dumps({"error": f"geometry mode: {e}"}))
            sys.exit(1)
        print(read_geometry(positional[0], page, scale, inks))
        sys.exit(0)

    # T57 (design §2.1): the colour mode — its own branch, so every other path stays byte-identical.
    if "--colours" in argv:
        try:
            page = int(argv[argv.index("--colours") + 1])
            positional = [a for i, a in enumerate(argv) if not a.startswith("--") and not (i > 0 and argv[i - 1] == "--colours")]
            if len(positional) < 1 or page < 1:
                raise ValueError("Usage: docling_bridge.py <doc_path> --colours <page>")
        except Exception as e:
            print(json.dumps({"error": f"colours mode: {e}"}))
            sys.exit(1)
        print(read_colours(positional[0], page))
        sys.exit(0)

    # T35 (§2.2): the page count alone, in milliseconds — the parser sizes the bridge's timeout by it.
    if "--page-count" in argv:
        positional = [a for a in argv if not a.startswith("--")]
        if len(positional) < 1:
            print(json.dumps({"error": "Usage: docling_bridge.py <doc_path> --page-count"}))
            sys.exit(1)
        print(page_count(positional[0]))
        sys.exit(0)

    # T35 (design D3): `--ocr auto | off` — the OCR mode for pages without a text layer (default auto).
    ocr_mode = "auto"
    if "--ocr" in argv:
        i = argv.index("--ocr")
        ocr_mode = argv[i + 1] if len(argv) > i + 1 else "auto"
    args = [a for i, a in enumerate(argv) if a != "--no-images" and a != "--ocr" and not (i > 0 and argv[i - 1] == "--ocr")]
    no_images = "--no-images" in sys.argv[1:]
    if len(args) < 1:
        print(json.dumps({"error": "Usage: docling_bridge.py <doc_path> [output_dir] [--no-images] [--ocr auto|off]"}))
        sys.exit(1)

    input_file = args[0]
    out_dir = args[1] if len(args) > 1 else None
    print(convert_document_to_markdown_pages(input_file, out_dir, no_images, ocr_mode))
