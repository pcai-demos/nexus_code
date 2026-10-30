#!/usr/bin/env python3
"""T34 — the .docx fixture for I-SECTION (tests/fixtures/golden/section_headings.docx).

A small Word file with the shapes the section anchor must handle: three heading levels, a heading REPEATED under two
parents ("Executive Summary" under "Report A" and under "Report B"), a blank heading paragraph, an ampersand in a
heading (docling escapes it — F-146), a table, and a sentinel token per section so a retrieval can be judged. Written
with python-docx (the manifest's independent docx engine); committed as bytes — regenerate only on purpose:

    python3 scripts/make_docx_fixture.py

The tokens are deliberate nonsense (rare by construction, never in the corpus): AMBERQUILL (Report A › Executive
Summary), COBALTREED (Report B › Executive Summary), SAFFRONVANE (Section 1: Methodology & Assumptions › 1.1 Scope).
"""
import os
import sys

try:
    import docx
except ImportError:
    sys.exit("python-docx is required: pip install python-docx")

OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "tests", "fixtures", "golden", "section_headings.docx")

d = docx.Document()
d.add_paragraph("Front matter before any heading: this paragraph has no section and stays that way.")
d.add_heading("Report A", level=1)
d.add_paragraph("Report A opens with a short introduction of its scope and its reader.")
d.add_heading("Executive Summary", level=2)
d.add_paragraph("Report A's summary states the AMBERQUILL finding: the first platform's three-year cost is the lower of the two under the stated assumptions, with the caveats listed in Section 1.")
d.add_heading("Section 1: Methodology & Assumptions", level=2)
d.add_paragraph("The methodology section explains how the numbers were derived and what was assumed about volumes, prices and staffing.")
d.add_heading("1.1 Scope", level=3)
d.add_paragraph("The scope paragraph carries the SAFFRONVANE boundary: on-premises and cloud, three years, one region, prices as of the stated month.")
t = d.add_table(rows=3, cols=2)
t.style = "Table Grid"
for r, (a, b) in enumerate([("Item", "Value"), ("Egress per GB", "0.0018"), ("Storage per GB-month", "0.020")]):
    t.cell(r, 0).text = a
    t.cell(r, 1).text = b
d.add_heading("", level=2)  # a blank heading paragraph (the TCO file has twenty of them)
d.add_paragraph("A paragraph after a blank heading still belongs to 1.1 Scope: the blank is skipped, never a section of its own.")
d.add_heading("Report B", level=1)
d.add_heading("Executive Summary", level=2)
d.add_paragraph("Report B's summary states the COBALTREED finding: the second comparison reverses the first under a different utilisation, and says why.")
d.add_paragraph("A closing paragraph of Report B's summary, long enough to be a chunk of its own when the splitter runs at a small chunk size, so that the section path is tested on a second chunk of the same section too.")
d.save(OUT)
print(f"wrote {OUT} ({os.path.getsize(OUT)} bytes)")
